"""Rapprochement d'un relevé Mobile Money ou bancaire avec les loyers attendus.

1. `lire_releve` transforme le fichier (CSV ou Excel) en lignes : date, montant, téléphone, libellé, référence.
2. `suggerer` propose pour chaque crédit le bail le plus probable (téléphone, montant, nom du locataire).
3. `encaisser_lignes` enregistre les loyers des lignes validées, en répartissant le montant sur les mois dus,
   et mémorise chaque ligne (empreinte) pour qu'un relevé importé deux fois ne soit jamais encaissé deux fois.
"""
import csv
import hashlib
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

from apps.core.exceptions import RegleMetier
from apps.core.utils import fcfa, journaliser
from apps.locatif import services
from apps.locatif.models import Contrat, Encaissement

from .models import LigneRapprochee

CHAMPS = ["date", "montant", "telephone", "libelle", "reference"]
MOTS_CLES = {
    "date": ["date", "jour", "horodatage", "time"],
    "montant": ["credit", "crédit", "montant", "amount", "somme", "valeur", "recu", "reçu"],
    "telephone": ["telephone", "téléphone", "tel", "phone", "msisdn", "numero", "numéro", "expediteur", "expéditeur", "emetteur", "émetteur", "from", "payeur"],
    "libelle": ["libelle", "libellé", "description", "motif", "detail", "détail", "nom", "name", "beneficiaire", "objet", "narration"],
    "reference": ["reference", "référence", "ref", "transaction", "id", "operation", "opération"],
}
LIMITE_LIGNES = 2000


def _sans_accents(t):
    return "".join(c for c in unicodedata.normalize("NFD", str(t or "").lower()) if unicodedata.category(c) != "Mn")


def _texte(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def lire_releve(fichier):
    """Renvoie (entêtes, lignes) d'un fichier CSV (séparateur détecté) ou Excel (.xlsx), lignes en texte."""
    nom = (getattr(fichier, "name", "") or "").lower()
    contenu = fichier.read()
    if nom.endswith(".xlsx") or contenu[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            feuille = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise RegleMetier("Fichier Excel illisible.")
        lignes = [[_texte(v) for v in r] for r in feuille.iter_rows(values_only=True)]
    else:
        texte = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                texte = contenu.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialecte = csv.Sniffer().sniff(texte[:4000], delimiters=";,\t|")
        except csv.Error:
            dialecte = csv.excel
            dialecte.delimiter = ";" if texte.count(";") > texte.count(",") else ","
        lignes = [[c.strip() for c in r] for r in csv.reader(io.StringIO(texte), dialecte)]
    lignes = [r for r in lignes if any(c for c in r)]
    if len(lignes) < 2:
        raise RegleMetier("Le relevé est vide ou ne contient qu'une ligne d'entêtes.")
    # La ligne d'entêtes est la première qui ressemble à des titres (relevés bancaires : quelques lignes d'infos avant)
    debut = next((i for i, r in enumerate(lignes[:15]) if sum(any(m in _sans_accents(c) for ms in MOTS_CLES.values() for m in ms) for c in r) >= 2), 0)
    entetes = [c or f"Colonne {i + 1}" for i, c in enumerate(lignes[debut])]
    corps = [r + [""] * (len(entetes) - len(r)) for r in lignes[debut + 1:debut + 1 + LIMITE_LIGNES]]
    return entetes, corps


def deviner_colonnes(entetes):
    """Colonne la plus probable pour chaque champ (index ou None)."""
    choix, pris = {}, set()
    for champ in CHAMPS:
        for i, e in enumerate(entetes):
            t = _sans_accents(e)
            if i not in pris and any(_sans_accents(m) in t for m in MOTS_CLES[champ]):
                if champ == "montant" and ("debit" in t or "débit" in t):
                    continue
                choix[champ] = i
                pris.add(i)
                break
        choix.setdefault(champ, None)
    return choix


def montant(v):
    t = re.sub(r"[^\d,.\-]", "", str(v or ""))
    if not t:
        return None
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".") if len(t.split(",")[-1]) <= 2 else t.replace(",", "")
    elif t.count(".") > 1 or (t.count(".") == 1 and len(t.split(".")[-1]) == 3):
        t = t.replace(".", "")
    try:
        return Decimal(t).quantize(Decimal(1))
    except InvalidOperation:
        return None


def date_releve(v):
    """Date du relevé, l'heure éventuelle étant ignorée (« 05/10/2026 14:32 », « 2026-10-05T14:32:00 »…)."""
    jour = re.split(r"[ T]", str(v or "").strip())[0]
    for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(jour, f).date()
        except ValueError:
            continue
    return None


def _chiffres(tel):
    d = re.sub(r"\D", "", tel or "")
    return d[-10:] if len(d) >= 10 else d


def empreinte(l):
    cle = f"{l['date']}|{l['montant']}|{_chiffres(l['telephone'])}|{(l['reference'] or l['libelle']).strip().lower()}"
    return hashlib.sha256(cle.encode()).hexdigest()[:40]


def _candidats():
    res = []
    for c in Contrat.objects.filter(statut_contrat="actif").select_related("locataire", "maison"):
        dus = [p for p in services.echeancier(c, apres=1) if p.reste > 0]
        res.append({
            "contrat": c, "tels": {t for t in (_chiffres(c.locataire.telephone1), _chiffres(c.locataire.telephone2)) if len(t) >= 8},
            "mots": {m for m in _sans_accents(c.locataire.nom).split() if len(m) >= 3},
            "reste_ancien": dus[0].reste if dus else None, "loyer": c.loyer_mensuel, "dette": sum((p.reste for p in dus), Decimal(0)),
        })
    return res


def suggerer(entetes, corps, colonnes):
    """Lignes normalisées avec, pour chaque crédit, le bail suggéré et un niveau de confiance."""
    cands = _candidats()
    deja = set(LigneRapprochee.objects.values_list("empreinte", flat=True))
    lignes = []
    for n, r in enumerate(corps):
        val = lambda champ: r[colonnes[champ]] if colonnes.get(champ) is not None and colonnes[champ] < len(r) else ""  # noqa: E731
        l = {"n": n, "date": date_releve(val("date")), "montant": montant(val("montant")), "telephone": val("telephone"),
             "libelle": val("libelle"), "reference": val("reference")}
        if l["montant"] is None or l["montant"] <= 0 or l["date"] is None:
            l.update(etat="ignoree", raison="Débit, montant ou date illisible")
            lignes.append(l)
            continue
        l["empreinte"] = empreinte(l)
        if l["empreinte"] in deja:
            l.update(etat="deja", raison="Déjà encaissée lors d'un import précédent")
            lignes.append(l)
            continue
        tel, mots = _chiffres(l["telephone"] + " " + l["libelle"]), set(_sans_accents(l["libelle"]).split())
        meilleur, score_max, raisons = None, 0, []
        for c in cands:
            score, pourquoi = 0, []
            if any(t and t[-8:] in tel for t in c["tels"]):
                score += 3
                pourquoi.append("téléphone")
            if l["montant"] in (c["reste_ancien"], c["loyer"]) or (c["dette"] and l["montant"] == c["dette"]):
                score += 2
                pourquoi.append("montant")
            if len(c["mots"] & mots) >= 1:
                score += 1 + (len(c["mots"] & mots) >= 2)
                pourquoi.append("nom")
            if score > score_max:
                meilleur, score_max, raisons = c, score, pourquoi
        if meilleur and score_max >= 2:
            c = meilleur["contrat"]
            l.update(etat="certain" if score_max >= 5 else "probable", contrat=c.pk, raison=" + ".join(raisons),
                     suggestion=f"{c.locataire.nom} — {c.maison.designation}", dette=meilleur["dette"])
        else:
            l.update(etat="a_affecter", raison="Aucun locataire reconnu")
        lignes.append(l)
    return lignes


@transaction.atomic
def encaisser_lignes(lignes, mode, source, request):
    """Encaisse chaque ligne validée sur le bail choisi, du plus ancien mois dû au plus récent."""
    if mode not in (Encaissement.Mode.MOBILE_MONEY, Encaissement.Mode.VIREMENT):
        raise RegleMetier("Mode de paiement invalide.")
    resultats = []
    for l in lignes:
        c = Contrat.objects.filter(pk=l.get("contrat"), statut_contrat="actif").select_related("locataire").first()
        d = {"n": l.get("n"), "montant": l.get("montant")}
        if c is None:
            resultats.append({**d, "ok": False, "message": "Bail introuvable ou terminé."})
            continue
        try:
            restant, d_op = Decimal(str(l["montant"])), datetime.strptime(str(l["date"]), "%Y-%m-%d").date()
        except (KeyError, ValueError, InvalidOperation):
            resultats.append({**d, "ok": False, "message": "Ligne illisible."})
            continue
        cle = empreinte({"date": d_op, "montant": restant, "telephone": l.get("telephone", ""), "libelle": l.get("libelle", ""), "reference": l.get("reference", "")})
        if LigneRapprochee.objects.filter(empreinte=cle).exists():
            resultats.append({**d, "ok": False, "message": "Déjà encaissée."})
            continue
        recus, observation = [], f"Relevé {source} — {l.get('reference') or l.get('libelle') or ''}".strip(" —")[:250]
        quand = datetime.combine(d_op, datetime.now().time())
        from django.utils import timezone

        quand = timezone.make_aware(quand) if timezone.is_naive(quand) else quand
        for p in [p for p in services.echeancier(c, apres=24) if p.reste > 0]:
            if restant <= 0:
                break
            part = min(restant, p.reste)
            enc = services.enregistrer_encaissement(c, part, p.mois, p.annee, mode, request=request, date_operation=quand,
                                                    observation=observation, controler_dette=False)
            recus.append(enc)
            restant -= part
        if not recus:
            resultats.append({**d, "ok": False, "message": "Aucun loyer à régler sur ce bail."})
            continue
        LigneRapprochee.objects.create(empreinte=cle, source=source[:60], date_operation=d_op, montant=Decimal(str(l["montant"])),
                                       telephone=l.get("telephone", "")[:30], libelle=(l.get("libelle") or "")[:255],
                                       reference=(l.get("reference") or "")[:120], contrat=c, encaissement=recus[0], reste_non_affecte=restant)
        resultats.append({**d, "ok": True, "locataire": c.locataire.nom, "recus": [{"id": e.pk, "reference": e.reference_recu, "periode": e.periode_concernee} for e in recus],
                          "message": f"Trop-perçu de {fcfa(restant)} à traiter à part." if restant > 0 else ""})
    ok = [r for r in resultats if r["ok"]]
    journaliser(request, "Rapprochement", f"Relevé {source} : {len(ok)} ligne(s) encaissée(s) sur {len(resultats)}")
    return resultats
