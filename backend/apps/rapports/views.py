"""Tableau de bord, rapports, calendrier, recherche globale et exports."""
import io
from datetime import date, datetime, timedelta
from decimal import Decimal

from django.db.models import Count, Q, Sum
from django.db.models.functions import ExtractMonth, TruncDate
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response

from apps.core.models import Journal
from apps.core.permissions import EstAdmin
from apps.core.utils import (
    MOIS, aujourdhui, date_du_mois, debut_fin_mois, index_periode, journaliser, libelle_periode,
)
from apps.finances.models import CompteCourantBailleur, MouvementCaisse
from apps.finances.services import solde_caisse
from apps.locatif import services
from apps.locatif.models import ChargeLocative, Contrat, Encaissement, MouvementCaution, RevisionLoyer
from apps.patrimoine.models import Bailleur, Locataire, MandatGestion, Maison
from apps.relations.models import MessageLocataire, MessageVisiteur, Reservation


def _somme(qs, champ):
    return qs.aggregate(t=Sum(champ))["t"] or Decimal(0)


def _contrats_actifs():
    return Contrat.objects.filter(statut_contrat="actif").select_related("maison__bailleur", "locataire")


def _situation_mois(contrats, mois, annee):
    """Pour chaque contrat actif : loyer du mois, montant payé, reste et statut (échéancier du mois)."""
    j = aujourdhui()
    paye = {
        l["contrat"]: l["t"]
        for l in Encaissement.objects.filter(periode_mois=mois, periode_annee=annee).values("contrat").annotate(t=Sum("montant_recu"))
    }
    lignes = []
    for c in contrats:
        debut = c.debut_facturation or c.date_debut
        if index_periode(mois, annee) < index_periode(debut.month, debut.year):
            continue
        loyer = services.loyer_pour_periode(c, mois, annee)
        p = Decimal(paye.get(c.pk) or 0)
        echeance = date_du_mois(annee, mois, c.date_debut.day)
        if p >= loyer:
            statut = "paye"
        elif echeance < j:
            statut = "en_retard" if p == 0 else "partiel_retard"
        else:
            statut = "partiel" if p > 0 else "a_venir"
        lignes.append({
            "contrat": c.pk, "reference": c.reference, "locataire": c.locataire.nom, "locataire_id": c.locataire_id,
            "maison": c.maison.designation, "loyer": loyer, "paye": p, "reste": max(Decimal(0), loyer - p),
            "echeance": echeance, "statut": statut,
        })
    return sorted(lignes, key=lambda l: (l["echeance"], l["locataire"]))


@api_view(["GET"])
def tableau_de_bord(request):
    j = aujourdhui()
    annee, mois = j.year, j.month
    mois_prec_d = (j.replace(day=1) - timedelta(days=1))
    enc = Encaissement.objects.all()
    actifs = list(_contrats_actifs())

    encaisse_annee = _somme(enc.filter(date_encaissement__year=annee), "montant_recu")
    encaisse_mois = _somme(enc.filter(date_encaissement__year=annee, date_encaissement__month=mois), "montant_recu")
    encaisse_mois_prec = _somme(
        enc.filter(date_encaissement__year=mois_prec_d.year, date_encaissement__month=mois_prec_d.month), "montant_recu"
    )

    situation = _situation_mois(actifs, mois, annee)
    attendu_mois = sum((l["loyer"] for l in situation), Decimal(0))
    collecte_mois = sum((min(l["paye"], l["loyer"]) for l in situation), Decimal(0))

    impayes_total, contrats_en_retard = Decimal(0), 0
    tranches = [("0–30 j", 30), ("31–60 j", 60), ("61–90 j", 90), ("Plus de 90 j", None)]
    anciennete = {t: Decimal(0) for t, _ in tranches}
    for c in actifs:
        imp = services.periodes_impayees(c)
        if imp:
            contrats_en_retard += 1
            impayes_total += sum((p.reste for p in imp), Decimal(0))
        for p in imp:
            jours = (j - date_du_mois(p.annee, p.mois, c.date_debut.day)).days
            tranche = next(t for t, borne in tranches if borne is None or jours <= borne)
            anciennete[tranche] += p.reste

    total_maisons = Maison.objects.count()
    occupees = Maison.objects.filter(statut="occupe").count()

    def par_mois(a):
        lignes = (
            enc.filter(date_encaissement__year=a).annotate(m=ExtractMonth("date_encaissement"))
            .values("m").annotate(t=Sum("montant_recu"))
        )
        d = {l["m"]: l["t"] for l in lignes}
        return [d.get(i, 0) for i in range(1, 13)]

    modes = list(
        enc.filter(date_encaissement__year=annee).values("mode_paiement")
        .annotate(nb=Count("id"), total=Sum("montant_recu")).order_by("-total")
    )
    libelles_modes = dict(Encaissement.Mode.choices)
    for m in modes:
        m["libelle"] = libelles_modes.get(m["mode_paiement"], m["mode_paiement"])

    # Recouvrement des 12 derniers mois : loyers attendus et encaissés pour chaque période
    tous = list(Contrat.objects.select_related("maison__bailleur", "locataire"))
    recouvrement = []
    for k in range(11, -1, -1):
        i = index_periode(mois, annee) - k
        m, a = i % 12 + 1, i // 12
        debut_m, fin_m = date(a, m, 1), date_du_mois(a, m, 31)
        en_cours = [c for c in tous if c.date_debut <= fin_m and (c.date_resiliation is None or c.date_resiliation >= debut_m)]
        lignes_m = _situation_mois(en_cours, m, a)
        recouvrement.append({
            "mois": m, "annee": a,
            "attendu": sum((l["loyer"] for l in lignes_m), Decimal(0)),
            "encaisse": sum((min(l["paye"], l["loyer"]) for l in lignes_m), Decimal(0)),
        })

    commissions = (
        MouvementCaisse.objects.filter(type_mouvement="Commission Gestion", date_operation__year=annee)
        .annotate(m=ExtractMonth("date_operation")).values("m").annotate(t=Sum("montant"))
    )
    commissions_mois = {l["m"]: l["t"] for l in commissions}

    occupation = {}
    for t, st in Maison.objects.values_list("type_maison", "statut"):
        ligne = occupation.setdefault(t or "Autre", {"type": t or "Autre", "occupees": 0, "libres": 0})
        ligne["occupees" if st == "occupe" else "libres"] += 1

    top_maisons = list(
        enc.filter(date_encaissement__year=annee).values("contrat__maison__designation")
        .annotate(total=Sum("montant_recu")).order_by("-total")[:5]
    )
    prochaines = sorted(
        [c for c in actifs if c.date_prochain_loyer and c.date_prochain_loyer >= j], key=lambda c: c.date_prochain_loyer
    )[:8]
    fins = [c for c in actifs if c.date_fin and j <= c.date_fin <= j + timedelta(days=30)]
    mandats_fin = MandatGestion.objects.en_vigueur().filter(
        duree="determinee", tacite_reconduction=False, date_fin__lte=j + timedelta(days=90)
    ).select_related("bailleur")

    derniers = enc.select_related("contrat__locataire", "contrat__maison")[:8]
    data = {
        "date": j,
        "kpi": {
            "encaisse_annee": encaisse_annee, "encaisse_mois": encaisse_mois, "encaisse_mois_precedent": encaisse_mois_prec,
            "attendu_mois": attendu_mois, "collecte_mois": collecte_mois,
            "taux_recouvrement": round(float(collecte_mois / attendu_mois * 100), 1) if attendu_mois else 0,
            "total_maisons": total_maisons, "maisons_occupees": occupees, "maisons_libres": total_maisons - occupees,
            "taux_occupation": round(occupees / total_maisons * 100, 1) if total_maisons else 0,
            "locataires_actifs": len({c.locataire_id for c in actifs}), "contrats_actifs": len(actifs),
            "contrats_en_retard": contrats_en_retard, "impayes_total": impayes_total,
            "bailleurs": Bailleur.objects.count(),
            "solde_bailleurs": _somme(CompteCourantBailleur.objects.all(), "montant"),
            "cautions_detenues": _somme(Contrat.objects.filter(statut_contrat="actif"), "depot_garantie")
            + _somme(MouvementCaution.objects.filter(contrat__statut_contrat="actif"), "montant"),
            "solde_caisse": solde_caisse(),
            "commissions_annee": _somme(MouvementCaisse.objects.filter(type_mouvement="Commission Gestion", date_operation__year=annee), "montant"),
            "revenu_potentiel": _somme(Maison.objects.filter(statut="disponible"), "loyer"),
            "charges_annee": _somme(ChargeLocative.objects.filter(date_charge__year=annee), "montant"),
            "revisions_annee": RevisionLoyer.objects.filter(created_at__year=annee).count(),
        },
        "alertes": {
            "reservations": Reservation.objects.filter(statut="en_attente").count(),
            "messages": MessageVisiteur.objects.filter(lu=False).count(),
            "messages_locataires": MessageLocataire.objects.filter(expediteur="locataire", lu=False).count(),
            "fins_contrat": [
                {"contrat": c.pk, "locataire": c.locataire.nom, "maison": c.maison.designation, "date_fin": c.date_fin, "jours": (c.date_fin - j).days}
                for c in fins
            ],
            "sorties_prevues": [
                {"contrat": c.pk, "locataire": c.locataire.nom, "maison": c.maison.designation, "date": c.date_sortie_prevue, "jours": (c.date_sortie_prevue - j).days}
                for c in actifs if c.date_sortie_prevue and c.date_sortie_prevue <= j + timedelta(days=45)
            ],
            "mandats_expirant": [
                {"bailleur": m.bailleur.nom, "bailleur_id": m.bailleur_id, "date_fin": m.date_fin, "jours": (m.date_fin - j).days}
                for m in mandats_fin
            ],
            "revisions": len(services.revisions_a_echeance(30, j)),
        },
        "mois_courant": {"libelle": libelle_periode(mois, annee), "lignes": situation},
        "graphique": {"annee": annee, "courante": par_mois(annee), "precedente": par_mois(annee - 1)},
        "modes": modes,
        "graphes": {
            "recouvrement": recouvrement,
            "impayes_anciennete": [{"tranche": t, "montant": anciennete[t]} for t, _ in tranches],
            "commissions": [{"mois": i, "total": commissions_mois.get(i, 0)} for i in range(1, 13)],
            "occupation": sorted(occupation.values(), key=lambda o: -(o["occupees"] + o["libres"])),
        },
        "top_maisons": [{"maison": t["contrat__maison__designation"], "total": t["total"]} for t in top_maisons],
        "prochaines_echeances": [
            {"contrat": c.pk, "locataire": c.locataire.nom, "maison": c.maison.designation, "loyer": c.loyer_mensuel, "date": c.date_prochain_loyer}
            for c in prochaines
        ],
        "derniers_paiements": [
            {"id": e.pk, "reference": e.reference_recu, "locataire": e.contrat.locataire.nom, "maison": e.contrat.maison.designation,
             "montant": e.montant_recu, "date": e.date_encaissement, "periode": e.periode_concernee, "mode": e.get_mode_paiement_display()}
            for e in derniers
        ],
    }
    if request.user.role == "admin":
        data["journal"] = list(
            Journal.objects.select_related("utilisateur").values("id", "action", "details", "date_action", "utilisateur__nom_complet")[:8]
        )
    return Response(data)


@api_view(["GET"])
def mois_paiements(request):
    """Échéancier d'un mois donné pour tous les contrats actifs."""
    j = aujourdhui()
    mois = int(request.query_params.get("mois", j.month))
    annee = int(request.query_params.get("annee", j.year))
    return Response({"libelle": libelle_periode(mois, annee), "lignes": _situation_mois(_contrats_actifs(), mois, annee)})


# ─── Rapports ─────────────────────────────────────────────────────────────────
def _lignes_encaissements(qs):
    return [
        {"reference": e.reference_recu, "date": e.date_encaissement, "locataire": e.contrat.locataire.nom,
         "maison": e.contrat.maison.designation, "bailleur": e.contrat.maison.bailleur.nom, "periode": e.periode_concernee,
         "mode": e.get_mode_paiement_display(), "montant": e.montant_recu, "commission": e.commission, "net": e.net_bailleur}
        for e in qs.select_related("contrat__locataire", "contrat__maison__bailleur").order_by("date_encaissement")
    ]


def _lignes_caisse(qs):
    return [
        {"date": m.date_operation, "type": m.type_mouvement, "montant": m.montant, "commentaire": m.commentaire, "par": m.effectue_par}
        for m in qs.order_by("date_operation")
    ]


def _lignes_versements(qs):
    return [
        {"date": o.date_operation, "bailleur": o.bailleur.nom, "montant": -o.montant, "mode": o.mode, "commentaire": o.commentaire}
        for o in qs.select_related("bailleur").order_by("date_operation")
    ]


def _periode_rapport(du: date, au: date):
    enc = Encaissement.objects.filter(date_encaissement__date__gte=du, date_encaissement__date__lte=au)
    caisse = MouvementCaisse.objects.filter(date_operation__date__gte=du, date_operation__date__lte=au)
    vers = CompteCourantBailleur.objects.filter(type_operation="versement_bailleur", date_operation__date__gte=du, date_operation__date__lte=au)
    par_bailleur = list(
        enc.values("contrat__maison__bailleur__nom")
        .annotate(montant=Sum("montant_recu"), commission=Sum("commission"), net=Sum("net_bailleur"), nb=Count("id"))
        .order_by("contrat__maison__bailleur__nom")
    )
    par_mode = list(enc.values("mode_paiement").annotate(montant=Sum("montant_recu"), nb=Count("id")))
    libelles = dict(Encaissement.Mode.choices)
    return {
        "du": du, "au": au,
        "totaux": {
            "encaisse": _somme(enc, "montant_recu"), "commissions": _somme(enc, "commission"), "net_bailleurs": _somme(enc, "net_bailleur"),
            "nb_encaissements": enc.count(),
            "entrees_caisse": _somme(caisse.filter(montant__gt=0), "montant"),
            "depenses": -_somme(caisse.filter(montant__lt=0), "montant"),
            "versements_bailleurs": -_somme(vers, "montant"),
        },
        "encaissements": _lignes_encaissements(enc),
        "caisse": _lignes_caisse(caisse),
        "versements": _lignes_versements(vers),
        "par_bailleur": [
            {"bailleur": l["contrat__maison__bailleur__nom"], "montant": l["montant"], "commission": l["commission"], "net": l["net"], "nb": l["nb"]}
            for l in par_bailleur
        ],
        "par_mode": [{"mode": libelles.get(l["mode_paiement"], l["mode_paiement"]), "montant": l["montant"], "nb": l["nb"]} for l in par_mode],
    }


def _date(v, defaut):
    try:
        return datetime.strptime(v, "%Y-%m-%d").date() if v else defaut
    except ValueError:
        return defaut


@api_view(["GET"])
def rapport_journalier(request):
    d = _date(request.query_params.get("date"), aujourdhui())
    data = _periode_rapport(d, d)
    data["titre"] = f"Rapport journalier du {d:%d/%m/%Y}"
    return Response(data)


@api_view(["GET"])
def rapport_periode(request):
    j = aujourdhui()
    du = _date(request.query_params.get("du"), j.replace(day=1))
    au = _date(request.query_params.get("au"), j)
    data = _periode_rapport(du, au)
    data["titre"] = f"Rapport du {du:%d/%m/%Y} au {au:%d/%m/%Y}"
    return Response(data)


@api_view(["GET"])
def rapport_mensuel(request):
    j = aujourdhui()
    mois = int(request.query_params.get("mois", j.month))
    annee = int(request.query_params.get("annee", j.year))
    du, au = debut_fin_mois(annee, mois)
    data = _periode_rapport(du, au)
    data["titre"] = f"Rapport mensuel — {libelle_periode(mois, annee)}"
    situation = _situation_mois(_contrats_actifs(), mois, annee)
    data["recouvrement"] = {
        "attendu": sum((l["loyer"] for l in situation), Decimal(0)),
        "collecte": sum((min(l["paye"], l["loyer"]) for l in situation), Decimal(0)),
        "lignes": situation,
    }
    return Response(data)


@api_view(["GET"])
def rapport_annuel(request):
    annee = int(request.query_params.get("annee", aujourdhui().year))
    du, au = date(annee, 1, 1), date(annee, 12, 31)
    data = _periode_rapport(du, au)
    data.pop("encaissements")
    data["titre"] = f"Rapport annuel {annee}"
    enc = Encaissement.objects.filter(date_encaissement__year=annee)
    caisse = MouvementCaisse.objects.filter(date_operation__year=annee)
    lignes = []
    for m in range(1, 13):
        e = enc.filter(date_encaissement__month=m)
        c = caisse.filter(date_operation__month=m)
        lignes.append({
            "mois": MOIS[m - 1], "encaisse": _somme(e, "montant_recu"), "commissions": _somme(e, "commission"),
            "net": _somme(e, "net_bailleur"), "nb": e.count(), "depenses": -_somme(c.filter(montant__lt=0), "montant"),
        })
    data["par_mois"] = lignes
    data["contrats"] = {
        "nouveaux": Contrat.objects.filter(date_contrat__year=annee).count(),
        "resilies": Contrat.objects.filter(date_resiliation__year=annee).count(),
        "actifs": Contrat.objects.filter(statut_contrat="actif").count(),
    }
    data.pop("caisse")
    return Response(data)


@api_view(["GET"])
def etat_occupation(request):
    maisons = Maison.objects.select_related("bailleur").prefetch_related("contrats__locataire").order_by("bailleur__nom", "designation")
    if b := request.query_params.get("bailleur"):
        maisons = maisons.filter(bailleur=b)
    lignes = []
    for m in maisons:
        c = next((c for c in m.contrats.all() if c.statut_contrat == "actif"), None)
        lignes.append({
            "maison": m.designation, "type": m.type_maison, "adresse": m.adresse, "bailleur": m.bailleur.nom,
            "statut": m.get_statut_display(), "loyer": c.loyer_mensuel if c else m.loyer,
            "locataire": c.locataire.nom if c else "", "telephone": c.locataire.telephone1 if c else "",
            "depuis": c.date_debut if c else None,
        })
    occ = sum(1 for l in lignes if l["locataire"])
    return Response({
        "titre": "État d'occupation du patrimoine", "date": aujourdhui(), "lignes": lignes,
        "totaux": {
            "total": len(lignes), "occupees": occ, "libres": len(lignes) - occ,
            "taux": round(occ / len(lignes) * 100, 1) if lignes else 0,
            "loyers_occupes": sum((l["loyer"] for l in lignes if l["locataire"]), Decimal(0)),
            "loyers_libres": sum((l["loyer"] for l in lignes if not l["locataire"]), Decimal(0)),
        },
    })


@api_view(["GET"])
def liste_impayes(request):
    j = aujourdhui()
    lignes = []
    qs = _contrats_actifs()
    if b := request.query_params.get("bailleur"):
        qs = qs.filter(maison__bailleur=b)
    for c in qs:
        imp = services.periodes_impayees(c)
        if not imp:
            continue
        lignes.append({
            "contrat": c.pk, "reference": c.reference, "locataire": c.locataire.nom, "telephone": c.locataire.telephone1,
            "maison": c.maison.designation, "bailleur": c.maison.bailleur.nom, "loyer": c.loyer_mensuel,
            "nb_mois": len(imp), "periodes": ", ".join(p.libelle for p in imp), "montant_du": sum((p.reste for p in imp), Decimal(0)),
            "depuis": c.date_prochain_loyer, "jours_retard": (j - c.date_prochain_loyer).days if c.date_prochain_loyer else 0,
            "caution_disponible": services.caution_disponible(c),
        })
    lignes.sort(key=lambda l: -l["jours_retard"])
    return Response({
        "titre": "Liste des impayés", "date": j, "lignes": lignes,
        "totaux": {"contrats": len(lignes), "montant": sum((l["montant_du"] for l in lignes), Decimal(0))},
    })


@api_view(["GET"])
def releve_bailleur(request, pk):
    return Response(donnees_releve(get_object_or_404(Bailleur, pk=pk), request))


def donnees_releve(b, request):
    """Relevé du compte courant d'un bailleur sur une période (par défaut : le mois en cours)."""
    j = aujourdhui()
    du = _date(request.query_params.get("du"), j.replace(day=1))
    au = _date(request.query_params.get("au"), j)
    ops = CompteCourantBailleur.objects.filter(bailleur=b)
    initial = _somme(ops.filter(date_operation__date__lt=du), "montant")
    periode = ops.filter(date_operation__date__gte=du, date_operation__date__lte=au).order_by("date_operation", "id")
    solde, lignes = initial, []
    for o in periode:
        solde += o.montant
        lignes.append({"date": o.date_operation, "libelle": o.commentaire or o.get_type_operation_display(),
                       "credit": o.montant if o.montant > 0 else 0, "debit": -o.montant if o.montant < 0 else 0, "solde": solde})
    m = b.mandat_actif()
    return {
        "bailleur": {"id": b.pk, "nom": b.nom, "code": b.code_bailleur, "telephone": b.telephone1, "adresse": b.adresse,
                     "taux": m.taux_commission if m else None},
        "du": du, "au": au, "solde_initial": initial, "solde_final": solde, "lignes": lignes,
        "credits": sum((l["credit"] for l in lignes), Decimal(0)), "debits": sum((l["debit"] for l in lignes), Decimal(0)),
        "biens": list(b.maisons.values("designation", "statut", "loyer")),
    }


@api_view(["GET"])
def recap_fiscal(request, pk):
    return Response(donnees_recap_fiscal(get_object_or_404(Bailleur, pk=pk), request))


def _mois_loues(contrats, annee):
    """Nombre de mois de l'année couverts par au moins un bail, jusqu'à aujourd'hui (mois entamé = mois compté)."""
    mois = set()
    for c in contrats:
        debut = max(c.date_debut, date(annee, 1, 1))
        fin = min(c.date_resiliation or c.date_fin or date(annee, 12, 31), date(annee, 12, 31), aujourdhui())
        i = debut.year * 12 + debut.month - 1
        while i <= fin.year * 12 + fin.month - 1:
            mois.add(i % 12)
            i += 1
    return len(mois)


def donnees_recap_fiscal(b, request):
    """Récapitulatif annuel des revenus locatifs d'un bailleur, bien par bien (aide à la déclaration des revenus fonciers)."""
    try:
        annee = int(request.query_params.get("annee") or aujourdhui().year)
    except ValueError:
        annee = aujourdhui().year
    encs = Encaissement.objects.filter(contrat__maison__bailleur=b, date_encaissement__year=annee)
    biens = []
    for m in b.maisons.order_by("designation"):
        e = encs.filter(contrat__maison=m)
        t = e.aggregate(brut=Sum("montant_recu"), commission=Sum("commission"), net=Sum("net_bailleur"), n=Count("id"))
        contrats = list(m.contrats.filter(date_debut__lte=date(annee, 12, 31)).exclude(date_resiliation__lt=date(annee, 1, 1)))
        if not t["n"] and not contrats:
            continue
        biens.append({
            "designation": m.designation, "type": m.type_maison, "adresse": m.adresse, "loyer_actuel": m.loyer,
            "mois_loues": _mois_loues(contrats, annee), "quittances": t["n"],
            "locataires": ", ".join(sorted({c.locataire.nom for c in contrats})),
            "brut": t["brut"] or 0, "commission": t["commission"] or 0, "net": t["net"] or 0,
        })
    ops = CompteCourantBailleur.objects.filter(bailleur=b, date_operation__year=annee)
    ajustements = [
        {"date": o.date_operation, "libelle": o.commentaire or o.get_type_operation_display(), "montant": o.montant}
        for o in ops.filter(type_operation=CompteCourantBailleur.Type.AJUSTEMENT).order_by("date_operation")
    ]
    tot = lambda cle: sum((x[cle] for x in biens), Decimal(0))  # noqa: E731
    total_ajustements = sum((a["montant"] for a in ajustements), Decimal(0))
    m = b.mandat_actif()
    return {
        "annee": annee,
        "bailleur": {"nom": b.nom, "code": b.code_bailleur, "telephone": b.telephone1, "adresse": b.adresse, "cni": b.numero_cni,
                     "taux": m.taux_commission if m else None, "mandat": m.reference if m else None},
        "biens": biens, "ajustements": ajustements,
        "totaux": {
            "brut": tot("brut"), "commission": tot("commission"), "net": tot("net"), "ajustements": total_ajustements,
            "revenu_net": tot("net") + total_ajustements,
            "verse": -(ops.filter(type_operation=CompteCourantBailleur.Type.VERSEMENT).aggregate(t=Sum("montant"))["t"] or 0),
            "mois_loues": sum(x["mois_loues"] for x in biens),
        },
        "annees": sorted({d.year for d in Encaissement.objects.filter(contrat__maison__bailleur=b).dates("date_encaissement", "year")} | {aujourdhui().year}, reverse=True),
    }


@api_view(["GET"])
def activite_agents(request):
    """Activité de chaque membre de l'équipe sur la période : encaissements, baux, relances, visites, états des lieux.

    L'administrateur voit toute l'équipe ; un agent ne voit que sa propre ligne.
    """
    from apps.comptes.models import Utilisateur
    from apps.locatif.models import EtatDesLieux, Relance
    from apps.relations.models import Reservation

    j = aujourdhui()
    du = _date(request.query_params.get("du"), j.replace(day=1))
    au = _date(request.query_params.get("au"), j)
    equipe = Utilisateur.objects.filter(is_active=True).order_by("nom_complet")
    if getattr(request.user, "role", "") != "admin":
        equipe = equipe.filter(pk=request.user.pk)
    encs = Encaissement.objects.filter(date_encaissement__date__range=(du, au))
    lignes = []
    for u in equipe:
        e = encs.filter(enregistre_par=u).aggregate(n=Count("id"), total=Sum("montant_recu"), commissions=Sum("commission"))
        visites = Reservation.objects.filter(agent=u, date_visite__range=(du, au))
        lignes.append({
            "id": u.pk, "nom": u.nom_complet, "role": u.get_role_display() if hasattr(u, "get_role_display") else u.role,
            "encaissements": e["n"], "montant": e["total"] or 0, "commissions": e["commissions"] or 0,
            "baux": Contrat.objects.filter(cree_par=u, date_contrat__range=(du, au)).count(),
            "relances": Relance.objects.filter(effectuee_par=u, date_relance__range=(du, au)).count(),
            "visites": visites.filter(statut="termine").count(), "visites_prevues": visites.exclude(statut__in=["termine", "annule"]).count(),
            "visiteurs_interesses": visites.filter(interesse=True).count(),
            "etats_des_lieux": EtatDesLieux.objects.filter(etabli_par=u, date_etat__range=(du, au)).count(),
        })
    sans_auteur = encs.filter(enregistre_par__isnull=True).aggregate(n=Count("id"), total=Sum("montant_recu"))
    return Response({
        "du": du, "au": au, "lignes": lignes,
        "automatiques": {"encaissements": sans_auteur["n"], "montant": sans_auteur["total"] or 0},
        "totaux": {k: sum(l[k] for l in lignes) for k in ("encaissements", "montant", "commissions", "baux", "relances", "visites", "etats_des_lieux")},
    })


# ─── Calendrier ───────────────────────────────────────────────────────────────
@api_view(["GET"])
def calendrier(request):
    j = aujourdhui()
    du = _date(request.query_params.get("du"), j.replace(day=1) - timedelta(days=7))
    au = _date(request.query_params.get("au"), du + timedelta(days=42))
    evts = []
    contrats = list(_contrats_actifs())
    i = index_periode(du.month, du.year)
    while i <= index_periode(au.month, au.year):
        mois, annee = i % 12 + 1, i // 12
        for l in _situation_mois(contrats, mois, annee):
            if du <= l["echeance"] <= au:
                evts.append({"id": f"loyer-{l['contrat']}-{i}", "type": "loyer", "date": l["echeance"], "statut": l["statut"],
                             "titre": l["locataire"], "detail": f"{l['maison']} — reste {int(l['reste']):,} FCFA".replace(",", " "),
                             "lien": f"/contrats/{l['contrat']}"})
        i += 1
    for r in Reservation.objects.filter(date_visite__range=(du, au), statut__in=["en_attente", "confirme"]).select_related("maison"):
        evts.append({"id": f"visite-{r.pk}", "type": "visite", "date": r.date_visite, "statut": r.statut,
                     "titre": f"Visite — {r.nom_visiteur}", "detail": f"{r.maison.designation} — {r.tel_visiteur}", "lien": "/reservations"})
    for rv in services.revisions_a_echeance(max(0, (au - j).days), j):
        if du <= rv["echeance"] <= au:
            evts.append({"id": f"revision-{rv['contrat']}", "type": "fin_contrat", "date": rv["echeance"], "statut": "info",
                         "titre": f"Révision triennale possible — {rv['locataire']}", "detail": f"{rv['maison']} — loyer {int(rv['loyer']):,} FCFA".replace(",", " "),
                         "lien": f"/contrats/{rv['contrat']}"})
    for c in Contrat.objects.filter(statut_contrat="actif", date_sortie_prevue__range=(du, au)).select_related("locataire", "maison"):
        evts.append({"id": f"sortie-{c.pk}", "type": "fin_contrat", "date": c.date_sortie_prevue, "statut": "info",
                     "titre": f"Sortie prévue — {c.locataire.nom}", "detail": f"{c.maison.designation} (préavis du {c.preavis_le:%d/%m/%Y})", "lien": f"/contrats/{c.pk}"})
    for c in Contrat.objects.filter(statut_contrat="actif", date_fin__range=(du, au)).select_related("locataire", "maison"):
        evts.append({"id": f"fin-{c.pk}", "type": "fin_contrat", "date": c.date_fin, "statut": "info",
                     "titre": f"Fin de bail — {c.locataire.nom}", "detail": c.maison.designation, "lien": f"/contrats/{c.pk}"})
    # Les mandats reconduits tacitement n'ont pas de fin : seuls les mandats à durée déterminée non reconduits apparaissent
    for m in MandatGestion.objects.en_vigueur().filter(duree="determinee", tacite_reconduction=False, date_fin__range=(du, au)).select_related("bailleur"):
        evts.append({"id": f"mandat-{m.pk}", "type": "fin_mandat", "date": m.date_fin, "statut": "info",
                     "titre": f"Fin de mandat — {m.bailleur.nom}", "detail": f"Taux {m.taux_commission}%", "lien": f"/bailleurs/{m.bailleur_id}"})
    return Response(sorted(evts, key=lambda e: e["date"]))


# ─── Recherche globale ────────────────────────────────────────────────────────
@api_view(["GET"])
def recherche(request):
    q = (request.query_params.get("q") or "").strip()
    if len(q) < 2:
        return Response([])
    res = []
    for b in Bailleur.objects.filter(Q(nom__icontains=q) | Q(code_bailleur__icontains=q) | Q(telephone1__icontains=q))[:5]:
        res.append({"type": "bailleur", "label": b.nom, "sub": b.code_bailleur, "lien": f"/bailleurs/{b.pk}"})
    for l in Locataire.objects.filter(Q(nom__icontains=q) | Q(telephone1__icontains=q))[:5]:
        res.append({"type": "locataire", "label": l.nom, "sub": l.telephone1, "lien": f"/locataires/{l.pk}"})
    for m in Maison.objects.filter(Q(designation__icontains=q) | Q(adresse__icontains=q))[:5]:
        res.append({"type": "maison", "label": m.designation, "sub": m.adresse[:60], "lien": f"/maisons/{m.pk}"})
    for c in Contrat.objects.filter(Q(locataire__nom__icontains=q) | Q(maison__designation__icontains=q)).select_related("locataire", "maison")[:5]:
        res.append({"type": "contrat", "label": f"{c.reference} — {c.maison.designation}", "sub": c.locataire.nom, "lien": f"/contrats/{c.pk}"})
    for e in Encaissement.objects.filter(reference_recu__icontains=q).select_related("contrat__locataire")[:5]:
        res.append({"type": "quittance", "label": e.reference_recu, "sub": f"{e.contrat.locataire.nom} — {e.periode_concernee}", "lien": f"/impression/quittance/{e.pk}"})
    return Response(res)


# ─── Export Excel générique ───────────────────────────────────────────────────
@api_view(["POST"])
def export_excel(request):
    """Reçoit {titre, colonnes:[{cle, libelle, type}], lignes:[{...}], totaux?} et renvoie un .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from apps.core.models import Parametres

    titre = (request.data.get("titre") or "Export")[:120]
    colonnes = request.data.get("colonnes") or []
    lignes = request.data.get("lignes") or []
    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31].replace("/", "-")
    ws.append([Parametres.charger().nom_entreprise])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([titre])
    ws["A2"].font = Font(size=11, color="5B6577")
    ws.append([f"Édité le {aujourdhui():%d/%m/%Y}"])
    ws.append([])
    ws.append([c.get("libelle", c.get("cle")) for c in colonnes])
    entete = ws.max_row
    for cell in ws[entete]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E2A47")
        cell.alignment = Alignment(vertical="center")
    for l in lignes[:50000]:
        row = []
        for c in colonnes:
            v = l.get(c["cle"])
            if c.get("type") == "montant" and v not in (None, ""):
                try:
                    v = float(v)
                except (TypeError, ValueError):
                    pass
            row.append(v)
        ws.append(row)
    for idx, c in enumerate(colonnes, start=1):
        lettre = get_column_letter(idx)
        if c.get("type") == "montant":
            for cell in ws[lettre][entete:]:
                cell.number_format = '#,##0" FCFA"'
        largeur = max([len(str(c.get("libelle", "")))] + [len(str(l.get(c["cle"], "") or "")) for l in lignes[:300]])
        ws.column_dimensions[lettre].width = min(max(10, largeur + 2), 60)
    ws.freeze_panes = ws.cell(row=entete + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    nom = "".join(ch if ch.isalnum() else "_" for ch in titre)[:60] or "export"
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{nom}.xlsx"'
    journaliser(request, "Export Excel", titre)
    return resp
